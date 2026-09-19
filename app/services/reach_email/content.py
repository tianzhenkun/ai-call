"""邮件名单与变量契约；不执行工作簿公式或模板表达式。"""

from __future__ import annotations

import asyncio
import re
from html import escape, unescape
from io import BytesIO
from zipfile import BadZipFile, ZipFile

import bleach
from dns.resolver import NXDOMAIN, NoAnswer
from email_validator import (
    EmailNotValidError,
    EmailUndeliverableError,
    caching_resolver,
    validate_email,
)
from openpyxl import Workbook, load_workbook

COLUMNS = ["客户姓名", "邮箱", "企业名称", "职务", "需求描述"]
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
VARIABLE = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def has_email_content(content: dict) -> bool:
    text = unescape(bleach.clean(content.get("html", ""), tags=[], strip=True)).strip()
    return bool(content.get("subject", "").strip() or text)


def email_address(value: str) -> str:
    value = value.strip()
    try:
        result = validate_email(
            value, check_deliverability=False, allow_smtputf8=False, strict=True
        )
    except EmailNotValidError as exc:
        raise ValueError("邮箱格式不正确") from exc
    return result.ascii_email.lower()


async def validate_recipient_domains(report: dict) -> None:
    """每个域名只查一次；解析失败或超时均不能视为校验通过。"""
    addresses = {
        r["values"]["邮箱"].rsplit("@", 1)[1]: r["values"]["邮箱"] for r in report["records"]
    }
    if not addresses:
        return
    retry = "邮箱域名暂时无法校验，请稍后重新上传重试"
    results = dict.fromkeys(addresses, retry)
    resolver = caching_resolver(timeout=2)
    pending = iter(addresses.items())

    def check(address):
        try:
            result = validate_email(
                address, dns_resolver=resolver, allow_smtputf8=False, strict=True
            )
        except EmailUndeliverableError as exc:
            if isinstance(exc.__cause__, NXDOMAIN):
                return "邮箱域名不存在，请检查 @ 后的拼写"
            if exc.__cause__ is None or isinstance(exc.__cause__, NoAnswer):
                return "邮箱域名未配置可用的收信服务，请核对邮箱地址"
            return retry
        except EmailNotValidError:
            return "邮箱格式不正确"
        return "" if getattr(result, "mx", None) else retry

    async def worker():
        for domain, address in pending:
            results[domain] = await asyncio.to_thread(check, address)

    try:
        await asyncio.wait_for(
            asyncio.gather(*(worker() for _ in range(min(16, len(addresses))))), 15
        )
    except asyncio.TimeoutError:
        # 保留已完成的结果，尚未完成的域名显示重试原因。
        pass
    records = []
    for record in report["records"]:
        email = record["values"]["邮箱"]
        reason = results[email.rsplit("@", 1)[1]]
        if reason:
            report["issues"].append({
                "row": record["row"],
                "email": record["originalEmail"],
                "reason": reason,
                "kind": "error",
            })
        else:
            records.append(record)
    for issue in report["issues"]:
        if issue["kind"] == "duplicate" and (
            reason := results.get(email_address(issue["email"]).rsplit("@", 1)[1])
        ):
            issue.update(kind="error", reason=reason)
            issue.pop("keptRow", None)
    report["issues"].sort(key=lambda issue: issue["row"])
    report["records"] = records
    report["validCount"] = len(records)
    report["errorCount"] = sum(i["kind"] == "error" for i in report["issues"])
    report["duplicateCount"] = len(report["issues"]) - report["errorCount"]
    report["canCreate"] = bool(records) and not report["errorCount"]


def clean_html(value: str) -> str:
    return bleach.clean(
        value,
        tags=[
            "p",
            "br",
            "div",
            "span",
            "strong",
            "b",
            "em",
            "i",
            "u",
            "s",
            "ul",
            "ol",
            "li",
            "a",
            "blockquote",
        ],
        attributes={"a": ["href", "title"]},
        protocols=["https", "http", "mailto"],
        strip=True,
    )


def validate_variables(value: str, columns: list[str]) -> None:
    missing = sorted(set(VARIABLE.findall(value)) - set(columns))
    if missing:
        raise ValueError("名单缺少变量列：" + "、".join(missing))
    remainder = VARIABLE.sub("", value)
    if "{{" in remainder or "}}" in remainder:
        raise ValueError("变量格式不正确")


def render_content(value: str, row: dict[str, str], *, html: bool = True) -> tuple[str, list[str]]:
    validate_variables(value, list(row))
    empty = sorted({key for key in VARIABLE.findall(value) if not row.get(key)})

    def replace(match):
        text = row.get(match.group(1), "")
        return escape(text, quote=True) if html else text

    return VARIABLE.sub(replace, value), empty


def xlsx_bytes(rows: list[list]) -> bytes:
    book = Workbook()
    for row in rows:
        # Explicit strings prevent spreadsheet formula injection in exported reports.
        book.active.append(row)
        for cell in book.active[book.active.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def parse_recipients(filename: str, payload: bytes) -> dict:
    if not filename.lower().endswith(".xlsx") or not 0 < len(payload) <= MAX_BYTES:
        raise ValueError("仅支持单个不超过 10 MB 的 XLSX 文件")
    try:
        with ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
                raise ValueError("工作簿解压后超过大小限制")
            if (
                any(e.flag_bits & 1 for e in entries)
                or "[Content_Types].xml" not in archive.namelist()
            ):
                raise ValueError("不支持加密或无效的工作簿")
        book = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ValueError("文件无法读取，请使用 XLSX 名单模板重新上传") from exc
    try:
        sheet = book.active
        if sheet.max_column is None or sheet.max_row is None:
            # 部分合法 XLSX 未写入可选尺寸信息，先从工作表计算实际范围。
            sheet.calculate_dimension(force=True)
        if sheet.max_column > 50 or sheet.max_row > MAX_ROWS + 1:
            raise ValueError("名单最多 10000 行、50 列")
        iterator = sheet.iter_rows()
        header = next(iterator, ())
        columns = [str(cell.value or "").strip() for cell in header]
        if (
            not columns
            or any(not c or len(c) > 64 or "{" in c or "}" in c for c in columns)
            or len(set(columns)) != len(columns)
        ):
            raise ValueError("表头不能为空、重复或包含变量括号")
        if "邮箱" not in columns:
            raise ValueError("名单缺少邮箱列")
        issues, records, seen = [], [], {}
        total = 0
        for number, cells in enumerate(iterator, 2):
            if all(cell.value is None for cell in cells):
                continue
            total += 1
            original_email = str(cells[columns.index("邮箱")].value or "")
            row = {key: str(cell.value or "").strip() for key, cell in zip(columns, cells, strict=False)}
            reason = ""
            if any(cell.data_type == "f" for cell in cells):
                reason = "名单不支持公式，请粘贴为文本"
            elif any(len(v) > 12000 for v in row.values()):
                reason = "单元格内容过长"
            else:
                try:
                    row["邮箱"] = email_address(row.get("邮箱", ""))
                except ValueError as exc:
                    reason = str(exc)
            if reason:
                issues.append(
                    {"row": number, "email": original_email, "reason": reason, "kind": "error"}
                )
            elif row["邮箱"] in seen:
                issues.append(
                    {
                        "row": number,
                        "email": original_email,
                        "reason": "重复邮箱，保留首次记录",
                        "kind": "duplicate",
                        "keptRow": seen[row["邮箱"]],
                    }
                )
            else:
                seen[row["邮箱"]] = number
                # 校验报告保留 Excel 原文，规范化地址仅用于去重和实际发信。
                records.append({"row": number, "values": row, "originalEmail": original_email})
        errors = sum(i["kind"] == "error" for i in issues)
        return {
            "columns": columns,
            "total": total,
            "validCount": len(records),
            "duplicateCount": len(issues) - errors,
            "errorCount": errors,
            "issues": issues,
            "records": records,
            "canCreate": bool(records) and not errors,
        }
    finally:
        book.close()
