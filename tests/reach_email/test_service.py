import asyncio
import json
from io import BytesIO

import pytest
from cryptography.fernet import Fernet
from openpyxl import Workbook
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.services.reach_email.models import TABLES, AIJob, Task
from app.services.reach_email.schema import AIInput, ContentInput, TaskInput
from app.services.reach_email.service import EmailService


def test_import_checks_domain_dns_and_reports_each_original_row(monkeypatch):
    import dns.exception
    import dns.resolver
    import dns.rrset
    from fastapi import UploadFile
    from openpyxl import load_workbook

    from app.api.v1.reach_email.controller import import_file, import_report
    from app.services.reach_email.content import xlsx_bytes

    calls = []

    def resolve(self, domain, kind, **kwargs):
        calls.append((domain, kind))
        if domain in ("qq.comd", "qq.xn--com-5w2h"):
            raise dns.resolver.NXDOMAIN()
        if domain == "broken.com":
            raise dns.exception.DNSException("resolver unavailable")
        if domain == "timeout.com":
            raise dns.exception.Timeout()
        if domain == "nomail.com":
            return dns.rrset.from_text(domain, 60, "IN", "MX", "0 .")
        return dns.rrset.from_text(domain, 60, "IN", "MX", "10 mx.qq.com.")

    monkeypatch.setattr(dns.resolver.Resolver, "resolve", resolve)

    async def run():
        engine = create_async_engine("sqlite+aiosqlite://")
        async with engine.begin() as conn:
            for table in TABLES:
                await conn.run_sync(lambda sync, t=table: t.create(sync))
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        async with sessions() as db:
            service = EmailService(db, "tenant-a", "1", "")
            response = await import_file(
                UploadFile(
                    filename="名单.xlsx",
                    file=BytesIO(
                        xlsx_bytes([
                            ["邮箱"],
                            ["745209176@qq.comd"],
                            ["second@qq.comd"],
                            ["745209176@qq.comd"],
                            ["a@nomail.com"],
                            ["a@timeout.com"],
                            ["a@qq.com"],
                            ["b@qq.com"],
                            ["991326583@qq.com的"],
                            ["991326583@qq.com的"],
                            ["a@broken.com"],
                        ])
                    ),
                ),
                service,
            )
            result = json.loads(response.body)["data"]
            assert result["errorCount"] == 8
            assert result["validCount"] == 2
            assert result["duplicateCount"] == 0
            assert not result["canCreate"]
            assert [issue["row"] for issue in result["issues"]] == [2, 3, 4, 5, 6, 9, 10, 11]
            assert result["issues"][0]["reason"] == "邮箱域名不存在，请检查 @ 后的拼写"
            assert result["issues"][3]["reason"] == "邮箱域名未配置可用的收信服务，请核对邮箱地址"
            assert "重试" in result["issues"][4]["reason"]
            assert result["issues"][5]["email"] == "991326583@qq.com的"
            assert result["issues"][6]["email"] == "991326583@qq.com的"
            assert "重试" in result["issues"][-1]["reason"]
            assert calls.count(("qq.comd", "MX")) == 1
            assert calls.count(("qq.com", "MX")) == 1
            assert calls.count(("qq.xn--com-5w2h", "MX")) == 1
            exported = await import_report(result["id"], service)
            book = load_workbook(BytesIO(exported.body), read_only=True)
            exported_rows = list(book.active.values)
            book.close()
            assert exported_rows[6][1] == "991326583@qq.com的"
            assert exported_rows[7][1] == "991326583@qq.com的"
        await engine.dispose()

    asyncio.run(run())


def test_task_identity_versions_and_missing_columns(valid_email_dns, task_settings):
    async def run():
        engine = create_async_engine("sqlite+aiosqlite://")
        async with engine.begin() as conn:
            for table in TABLES:
                await conn.run_sync(lambda sync, t=table: t.create(sync))
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        async with sessions() as db:
            a = EmailService(db, "tenant-a", "1", Fernet.generate_key().decode())
            book = Workbook()
            book.active.append(["邮箱", "客户姓名"])
            book.active.append(["a@example.com", ""])
            buf = BytesIO()
            book.save(buf)
            imp = await a.import_file("test.xlsx", buf.getvalue())
            task = await a.create_task(TaskInput(name="任务", importId=imp["id"], settings=task_settings))
            # 旧任务缺少公司信息时仍可生成草稿，执行时另行校验完整性。
            row = await db.get(Task, task["id"])
            row.settings = "{}"
            blank = await a.ai_job(task["id"], AIInput(version=1, action="generate"))
            blank_payload = json.loads((await db.get(AIJob, blank["id"])).payload)
            assert blank_payload["subject"] == blank_payload["content"] == blank_payload["instruction"] == ""
            row = await db.get(Task, task["id"])
            row.settings = json.dumps({"companyName": "发件方", "companyDescription": "提供产品演示服务"})
            generated = await a.ai_job(task["id"], AIInput(version=1, action="generate"))
            payload = json.loads((await db.get(AIJob, generated["id"])).payload)
            assert payload["context"]["companyName"] == "发件方"
            assert payload["allowed_variables"] == ["邮箱", "客户姓名"]
            with pytest.raises(ValueError, match="不存在"):
                await EmailService(db, "tenant-b", "1", "").task(task["id"])
            with pytest.raises(ValueError, match="不存在"):
                await EmailService(db, "tenant-a", "2", "").task(task["id"])
            changed = await a.save_content(
                task["id"], ContentInput(version=1, subject="你好", html="<p>{{客户姓名}}</p>", signatureName="商务签名")
            )
            assert changed["version"] == 2
            assert changed["content"]["signatureName"] == "商务签名"
            row = await db.get(Task, task["id"])
            saved_at = (await a.task_result(row))["contentUpdatedAt"]
            row.status = 'ended'
            await db.flush()
            assert (await a.task_result(row))["contentUpdatedAt"] == saved_at
            row.status = 'unstarted'
            with pytest.raises(ValueError, match="版本"):
                await a.save_content(task["id"], ContentInput(version=1, subject="覆盖", html="旧"))
            assert (await db.get(Task, task["id"])).name == "任务"
        await engine.dispose()

    asyncio.run(run())


def test_account_credentials_and_used_sender_are_protected():
    from app.services.reach_email.models import Account, Message, now
    from app.services.reach_email.schema import AccountInput
    from app.services.reach_email.security import CredentialCipher

    async def run():
        engine = create_async_engine("sqlite+aiosqlite://")
        async with engine.begin() as conn:
            for table in TABLES:
                await conn.run_sync(lambda sync, t=table: t.create(sync))
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        key = Fernet.generate_key().decode()
        async with sessions() as db:
            service = EmailService(db, "tenant", "owner", key)
            data = AccountInput(
                name="账号",
                email="sender@example.com",
                smtpHost="smtp.example.com",
                smtpUsername="sender",
                smtpPassword="private-secret",
                imapHost="imap.example.com",
                imapUsername="sender",
                imapPassword="other-private",
            )
            result = await service.save_account(data)
            assert "private-secret" not in str(result)
            account = await db.get(Account, result["id"])
            assert "private-secret" not in account.secret
            assert (
                CredentialCipher(key).decrypt(account.secret)["smtp_password"] == "private-secret"
            )
            account.smtp_status = account.imap_status = "ok"
            await service.save_account(
                data.model_copy(update={"smtp_password": "", "imap_password": ""}), account.id
            )
            assert account.smtp_status == "ok"
            service.add(
                Message,
                account_id=account.id,
                dedup_key="sent",
                message_id="<sent@example.com>",
                to_email="to@example.com",
                subject="subject",
                html="body",
                status="accepted",
                sent_at=now(),
            )
            await db.flush()
            with pytest.raises(ValueError, match="邮箱地址"):
                await service.save_account(
                    data.model_copy(update={"email": "new@example.com"}), account.id
                )
        await engine.dispose()

    asyncio.run(run())


def test_start_withdraw_stop_and_missing_column_are_atomic(valid_email_dns, task_settings):
    from datetime import timedelta, timezone

    from sqlalchemy import select

    from app.services.reach_email.models import Account, Lead, Message, now
    from app.services.reach_email.schema import StartInput

    async def run():
        engine = create_async_engine("sqlite+aiosqlite://")
        async with engine.begin() as conn:
            for table in TABLES:
                await conn.run_sync(lambda sync, t=table: t.create(sync))
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        async with sessions() as db:
            service = EmailService(db, "tenant", "owner", Fernet.generate_key().decode())
            service.add(
                Account,
                name="sender",
                email="s@example.com",
                config="{}",
                secret="encrypted",
                smtp_status="ok",
                imap_status="ok",
            )
            book = Workbook()
            book.active.append(["邮箱", "客户姓名", "需求描述"])
            book.active.append(["to@example.com", "", "说明"])
            buf = BytesIO()
            book.save(buf)
            imp = await service.import_file("recipients.xlsx", buf.getvalue())
            task = await service.create_task(TaskInput(name="t", importId=imp["id"], settings=task_settings))
            task = await service.save_content(
                task["id"], ContentInput(version=1, subject="hello", html="{{职务}}")
            )
            with pytest.raises(ValueError, match="名单缺少变量列"):
                await service.start(task["id"], StartInput(version=task["version"], requestId="x"))
            assert not (await db.scalars(select(Message))).all()
            task = await service.save_content(
                task["id"],
                ContentInput(
                    version=task["version"], subject="hello", html="{{客户姓名}} {{需求描述}}"
                ),
            )
            scheduled = now().replace(tzinfo=timezone.utc) + timedelta(hours=1)
            assert (await service.preflight(task["id"]))["recipientCount"] == 1
            result = await service.start(
                task["id"],
                StartInput(version=task["version"], requestId="start-1", scheduledAt=scheduled),
            )
            assert result["status"] == "scheduled"
            assert result["warnings"][0]["variables"] == ["客户姓名"]
            await service.start(
                task["id"],
                StartInput(version=task["version"], requestId="start-1", scheduledAt=scheduled),
            )
            assert len((await db.scalars(select(Message))).all()) == 1
            task = await service.transition(task["id"], result["version"], "withdraw")
            assert task["status"] == "unstarted"
            assert not (await db.scalars(select(Message))).all()
            lead = await db.scalar(select(Lead))
            assert lead.first_message_id is None
            result = await service.start(
                task["id"], StartInput(version=task["version"], requestId="start-2")
            )
            result = await service.transition(task["id"], result["version"], "stop")
            assert result["endedReason"] == "stopped"
            assert (await db.scalar(select(Message))).status == "cancelled"
            with pytest.raises(ValueError):
                await service.start(
                    task["id"], StartInput(version=result["version"], requestId="start-3")
                )
        await engine.dispose()

    asyncio.run(run())


def test_required_settings_and_empty_body_block_start_without_enqueuing(valid_email_dns, task_settings):
    from datetime import timedelta, timezone

    from sqlalchemy import select

    from app.services.reach_email.models import Lead, Message, now
    from app.services.reach_email.schema import StartInput

    async def run():
        engine = create_async_engine("sqlite+aiosqlite://")
        async with engine.begin() as conn:
            for table in TABLES:
                await conn.run_sync(lambda sync, t=table: t.create(sync))
        async with async_sessionmaker(engine, expire_on_commit=False)() as db:
            service = EmailService(db, "tenant", "owner", Fernet.generate_key().decode())
            book = Workbook()
            book.active.append(["邮箱"])
            book.active.append(["to@example.com"])
            buf = BytesIO()
            book.save(buf)
            imported = await service.import_file("recipients.xlsx", buf.getvalue())
            data = {"name": "任务", "importId": imported["id"]}
            with pytest.raises(ValueError):
                TaskInput(**data)
            task = await service.create_task(TaskInput(**data, settings=task_settings))
            row = await db.get(Task, task["id"])
            valid_content = {"subject": "交流邀请", "html": "<p>您好</p>"}
            cases = [({}, valid_content, "公司名称")]
            for field, value, error in (
                ("companyName", " \t", "公司名称"),
                ("companyWebsite", "", "公司官网"),
                ("companyDescription", "\u00a0", "公司简介"),
                ("companyWebsite", "example.com", "http 或 https"),
                ("companyWebsite", "https:example.com", "http 或 https"),
                ("companyWebsite", "ftp://example.com", "http 或 https"),
                ("dailyLimit", 0, "dailyLimit"),
                ("followUpIntervalDays", 0, "followUpIntervalDays"),
                ("followUpCount", 6, "followUpCount"),
            ):
                cases.append(({**task_settings, field: value}, valid_content, error))
            for settings, _, _ in cases:
                # 新建和修改使用相同入参校验，尚未进入数据库写入流程。
                for version in (None, 1):
                    with pytest.raises(ValueError):
                        TaskInput(**data, version=version, settings=settings)
            for html in ("", "<p><br></p>", "<p>&nbsp; &#160; </p>", "<p><strong> </strong></p>"):
                cases.append((task_settings, {"subject": "交流邀请", "html": html}, "邮件主题和正文"))
            cases.append((task_settings, {"subject": "  ", "html": "<p>您好</p>"}, "邮件主题和正文"))
            scheduled = now().replace(tzinfo=timezone.utc) + timedelta(hours=1)
            for settings, content, error in cases:
                # 模拟历史任务的不完整数据，预检和直接执行均须拦截。
                row.settings = json.dumps(settings)
                row.content = json.dumps(content)
                await db.flush()
                with pytest.raises(ValueError, match=error):
                    await service.preflight(row.id)
                for due in (None, scheduled):
                    with pytest.raises(ValueError, match=error):
                        await service.start(row.id, StartInput(version=1, requestId="blocked", scheduledAt=due))
                assert row.status == "unstarted" and row.version == 1
                assert row.start_key is None and row.scheduled_at is None
                assert not (await db.scalars(select(Message))).all()
                assert (await db.scalar(select(Lead))).first_message_id is None
        await engine.dispose()

    asyncio.run(run())


def test_manual_resolution_and_suppression_preserve_scope_and_never_resubmit():
    from pydantic import ValidationError
    from sqlalchemy import select

    from app.services.reach_email.models import Attempt, Lead, Message, Suppression
    from app.services.reach_email.schema import ResolveInput

    with pytest.raises(ValidationError):
        ResolveInput(outcome="accepted", note="       ")

    async def run():
        engine = create_async_engine("sqlite+aiosqlite://")
        async with engine.begin() as conn:
            for table in TABLES:
                await conn.run_sync(lambda sync, t=table: t.create(sync))
        sessions = async_sessionmaker(engine, expire_on_commit=False)
        async with sessions() as db:
            service = EmailService(db, "tenant", "owner", "")
            other = EmailService(db, "other-tenant", "owner", "")
            task = service.add(
                Task,
                name="task",
                import_id="test-import",
                status="running",
                settings='{"followUpEnabled":true,"followUpCount":2}',
            )
            await db.flush()
            lead = service.add(Lead, task_id=task.id, email="a@example.com", values="{}")
            await db.flush()
            msg = service.add(
                Message,
                task_id=task.id,
                lead_id=lead.id,
                status="unknown",
                dedup_key="first:test",
                to_email=lead.email,
                message_id="<first@example.com>",
                subject="s",
                html="body",
            )
            await db.flush()
            attempt = service.add(
                Attempt, task_id=task.id, message_id=msg.id, account_id="account", status="unknown"
            )
            pending = service.add(
                Message,
                task_id=task.id,
                lead_id=lead.id,
                status="queued",
                dedup_key="next:test",
                to_email=lead.email,
                message_id="<next@example.com>",
                subject="s",
                html="body",
            )
            foreign = other.add(
                Message,
                status="queued",
                dedup_key="other:test",
                to_email=lead.email,
                message_id="<other@example.com>",
                subject="s",
                html="body",
            )
            await db.flush()
            with pytest.raises(ValueError, match="不存在"):
                await other.resolve_message(
                    msg.id, ResolveInput(outcome="accepted", note="verified by provider")
                )
            result = await service.resolve_message(
                msg.id, ResolveInput(outcome="accepted", note=" verified by provider ")
            )
            assert result["resolutionNote"] == "verified by provider"
            assert msg.status == attempt.status == "accepted"
            assert lead.next_follow_up_at is not None
            with pytest.raises(ValueError, match="待核对"):
                await service.resolve_message(
                    msg.id, ResolveInput(outcome="failed", note="cannot confirm twice")
                )
            assert len((await db.scalars(select(Message))).all()) == 3
            await service.do_not_contact(lead.id, "manual")
            assert pending.status == "cancelled"
            assert foreign.status == "queued"
            assert lead.stopped_reason == "do_not_contact" and lead.next_follow_up_at is None
            assert len((await db.scalars(select(Suppression))).all()) == 1
            assert msg.status == "accepted"
        await engine.dispose()

    asyncio.run(run())


def test_secret_file_conflict_and_restart_decryption(tmp_path):
    from app.services.reach_email.security import CredentialCipher, configured_secret

    path = tmp_path / "email-key"
    key = Fernet.generate_key().decode()
    path.write_text(key + "\n")
    resolved = configured_secret("", str(path), "EMAIL_KEY")
    encrypted = CredentialCipher(resolved).encrypt({"smtp_password": "test-password"})
    assert CredentialCipher(configured_secret("", str(path), "EMAIL_KEY")).decrypt(encrypted) == {
        "smtp_password": "test-password"
    }
    with pytest.raises(ValueError, match="不能同时配置"):
        configured_secret(key, str(path), "EMAIL_KEY")
    with pytest.raises(ValueError, match="文件无法读取"):
        configured_secret("", str(tmp_path / "missing"), "EMAIL_KEY")
    with pytest.raises(ValueError, match="无法解密"):
        CredentialCipher(Fernet.generate_key().decode()).decrypt(encrypted)
